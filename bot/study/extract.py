from __future__ import annotations

import io
from pathlib import Path

import pypdf
from docx import Document
from pptx import Presentation

# Markers carry their own separator: "## p.7" but "## slide 7".
PAGE_MARKER = "p."
SLIDE_MARKER = "slide "
PART_MARKER = "part "

PDF_MIMES = ("application/pdf",)
PPTX_MIMES = ("application/vnd.openxmlformats-officedocument.presentationml.presentation",)
DOCX_MIMES = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",)
TEXT_MIMES = ("text/plain", "text/markdown", "text/x-markdown")
_TEXT_SUFFIXES = (".txt", ".md", ".markdown", ".rst", ".tex", ".csv", ".tsv", ".log")
CODE_MIMES = ("application/json", "application/x-python", "application/x-sh", "application/xml",
              "application/x-yaml", "application/javascript", "application/typescript")
_CODE_SUFFIXES = (
    ".py", ".ipynb", ".js", ".jsx", ".ts", ".tsx", ".java", ".kt", ".c", ".h", ".cpp", ".hpp", ".cc",
    ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".scala", ".lua", ".r", ".m", ".sh", ".bash", ".zsh",
    ".sql", ".html", ".css", ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".xml", ".env",
)
_IMAGE_SUFFIXES = (".jpg", ".jpeg", ".png", ".heic", ".webp", ".gif")
_DOCX_PARAGRAPHS_PER_CHUNK = 40


def extract_pdf(data: bytes) -> list[tuple[int, str]]:
    """(page number from 1, text) for every page; a page with no text is kept as ""."""
    reader = pypdf.PdfReader(io.BytesIO(data))
    return [(i, (page.extract_text() or "").strip()) for i, page in enumerate(reader.pages, 1)]


def extract_pptx(data: bytes) -> list[tuple[int, str]]:
    """(slide number from 1, text) — every text frame, then the speaker notes."""
    presentation = Presentation(io.BytesIO(data))
    slides: list[tuple[int, str]] = []
    for i, slide in enumerate(presentation.slides, 1):
        parts = [shape.text_frame.text.strip() for shape in slide.shapes if shape.has_text_frame]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")
        slides.append((i, "\n".join(p for p in parts if p)))
    return slides


def extract_docx(data: bytes) -> list[tuple[int, str]]:
    """(part number from 1, text) — paragraphs (headings included) chunked ~40 per part."""
    document = Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    chunks: list[tuple[int, str]] = []
    for i in range(0, len(paragraphs), _DOCX_PARAGRAPHS_PER_CHUNK):
        part = paragraphs[i : i + _DOCX_PARAGRAPHS_PER_CHUNK]
        chunks.append((i // _DOCX_PARAGRAPHS_PER_CHUNK + 1, "\n".join(part)))
    return chunks


def is_code(filename: str, mime: str) -> bool:
    return Path((filename or "").lower()).suffix in _CODE_SUFFIXES or (mime or "").lower() in CODE_MIMES


def is_plain_text(filename: str, mime: str) -> bool:
    """Anything that is its own text: notes, data, code. `text/*` mimes count whatever the name."""
    suffix = Path((filename or "").lower()).suffix
    mime = (mime or "").lower()
    return suffix in _TEXT_SUFFIXES or mime in TEXT_MIMES or mime.startswith("text/") or is_code(filename, mime)


def looks_like_text(data: bytes) -> bool:
    """An unknown file that decodes as UTF-8 with no control bytes is text, whatever it's called."""
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        return False
    return bool(text.strip()) and not any(ord(ch) < 32 and ch not in "\t\n\r" for ch in text)


def extract_text(data: bytes) -> str:
    """A .txt/.md upload is already the text; it goes in as written."""
    return data.decode("utf-8", errors="replace").strip()


def render_pages(pages: list[tuple[int, str]], marker: str) -> str:
    """One `## p.N` / `## slide N` heading per page, empty pages skipped."""
    blocks = [f"## {marker}{n}\n{text.strip()}" for n, text in pages if text.strip()]
    return "\n\n".join(blocks)


def guess_kind(filename: str, mime: str) -> str:
    name = (filename or "").lower()
    mime = (mime or "").lower()
    if name.endswith(".pptx") or mime in PPTX_MIMES:
        return "slides"
    if name.endswith(".pdf") or mime in PDF_MIMES:
        return "chapter"
    if is_code(name, mime):
        return "code"
    if name.endswith(".docx") or mime in DOCX_MIMES or is_plain_text(name, mime):
        return "notes"
    if mime.startswith("image/") or Path(name).suffix in _IMAGE_SUFFIXES:
        return "photo"
    return "other"
