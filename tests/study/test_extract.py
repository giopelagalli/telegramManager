import io

import pypdf
from pptx import Presentation

from bot.study.extract import (
    PAGE_MARKER,
    SLIDE_MARKER,
    extract_pdf,
    extract_pptx,
    guess_kind,
    render_pages,
)


class FakePage:
    def __init__(self, text):
        self._text = text

    def extract_text(self):
        return self._text


def test_extract_pdf_numbers_pages_and_keeps_empty_ones(monkeypatch):
    class FakeReader:
        def __init__(self, stream):
            self.pages = [FakePage(" One "), FakePage(None), FakePage("Three")]

    monkeypatch.setattr(pypdf, "PdfReader", FakeReader)
    assert extract_pdf(b"%PDF") == [(1, "One"), (2, ""), (3, "Three")]


def _deck() -> bytes:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[5])
    first.shapes.title.text = "Pointers"
    first.notes_slide.notes_text_frame.text = "Explain the stack"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_extract_pptx_reads_text_frames_and_notes():
    slides = extract_pptx(_deck())
    assert slides[0] == (1, "Pointers\nNotes: Explain the stack")
    assert slides[1] == (2, "")


def test_render_pages_skips_empty_pages():
    pages = [(1, "One"), (2, "  "), (3, "Three")]
    assert render_pages(pages, PAGE_MARKER) == "## p.1\nOne\n\n## p.3\nThree"
    assert render_pages(pages, SLIDE_MARKER) == "## slide 1\nOne\n\n## slide 3\nThree"
    assert render_pages([], PAGE_MARKER) == ""


def test_guess_kind():
    assert guess_kind("lecture7.pptx", "") == "slides"
    assert guess_kind("ch4.pdf", "application/pdf") == "chapter"
    assert guess_kind("", "application/pdf") == "chapter"
    assert guess_kind("board.jpg", "") == "photo"
    assert guess_kind("board", "image/jpeg") == "photo"
    assert guess_kind("data.zip", "application/zip") == "other"
